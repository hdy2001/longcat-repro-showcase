# -*- coding: utf-8 -*-
"""
LW 咨询风格 4 页 PPT 生成脚本
主题：AI 社交软件市场竞争格局与未来展望
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- 设计系统 ----------
RED      = RGBColor(0xD2, 0x26, 0x30)   # 主红
DARKRED  = RGBColor(0xA8, 0x1C, 0x24)   # 深红
DARK     = RGBColor(0x1F, 0x1F, 0x1F)   # 正文黑
GRAY     = RGBColor(0x55, 0x50, 0x4E)   # 正文灰
LGRAY    = RGBColor(0x8A, 0x83, 0x80)   # 脚注灰
BORDER   = RGBColor(0xE4, 0xDA, 0xD7)   # 卡片边框
BANNER_BG= RGBColor(0xFB, 0xE9, 0xE7)   # 底部横幅底色
FOOTER   = RGBColor(0x9A, 0x93, 0x90)
FONT     = "微软雅黑"

SW, SH = 13.333, 7.5
MX = 0.55                     # 页边距
CW = SW - 2 * MX              # 内容宽 12.233
CONTENT_Y = 1.62
CONTENT_B = 6.42
BANNER_Y, BANNER_H = 6.55, 0.62

prs = Presentation()
prs.slide_width  = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


def set_font(run, size=10, bold=False, color=DARK, name=FONT):
    f = run.font
    f.size = Pt(size); f.bold = bold; f.color.rgb = color; f.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set('typeface', name)


def txbox(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
          space_after=2, line_spacing=1.0, wrap=True):
    """lines: list of paragraphs; each paragraph = list of (text, size, bold, color) runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        for (t, sz, bd, c) in para:
            r = p.add_run(); r.text = t; set_font(r, sz, bd, c)
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, line_w=0.75, round_=False, adj=0.08):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if round_:
        try: shp.adjustments[0] = adj
        except Exception: pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def hline(slide, x, y, w, color=BORDER, h=0.012):
    return rect(slide, x, y, w, h, fill=color)


def badge(slide, x, y, num, d=0.30, fill=RED):
    """红色编号圆点"""
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = fill
    c.line.fill.background(); c.shadow.inherit = False
    tf = c.text_frame; tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(num); set_font(r, 10.5, True, RGBColor(0xFF, 0xFF, 0xFF))
    return c


def target_icon(slide, x, y, d=0.34):
    """底部横幅左侧的同心圆靶心图标"""
    o1 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    o1.fill.background(); o1.line.color.rgb = RED; o1.line.width = Pt(1.6); o1.shadow.inherit = False
    o2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + d*0.28), Inches(y + d*0.28),
                                Inches(d*0.44), Inches(d*0.44))
    o2.fill.solid(); o2.fill.fore_color.rgb = RED; o2.line.fill.background(); o2.shadow.inherit = False


def header(slide, page, title, subtitle):
    txbox(slide, MX, 0.30, 3.2, 0.3, [[(f"Page {page:02d}", 11.5, True, RED)]])
    rect(slide, MX, 0.60, 0.62, 0.035, fill=RED)
    txbox(slide, SW - MX - 1.5, 0.30, 1.5, 0.3, [[(f"{page:02d} / 04", 11, True, RED)]],
          align=PP_ALIGN.RIGHT)
    txbox(slide, MX, 0.74, CW, 0.5, [[(title, 24, True, DARK)]])
    txbox(slide, MX, 1.28, CW, 0.32, [[(subtitle, 11, False, GRAY)]])


def footer(slide, text):
    hline(slide, MX, 7.16, CW, color=BORDER)
    txbox(slide, MX, 7.22, CW, 0.24, [[(text, 7.5, False, LGRAY)]])


def banner(slide, text):
    rect(slide, MX, BANNER_Y, CW, BANNER_H, fill=BANNER_BG, round_=True, adj=0.28)
    rect(slide, MX, BANNER_Y, 0.07, BANNER_H, fill=RED)
    target_icon(slide, MX + 0.22, BANNER_Y + (BANNER_H - 0.34) / 2, d=0.34)
    txbox(slide, MX + 0.72, BANNER_Y, CW - 0.95, BANNER_H, [[(text, 11.5, True, DARKRED)]],
          anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)


def card(slide, x, y, w, h, title=None, title_color=DARK, border=BORDER):
    rect(slide, x, y, w, h, fill=RGBColor(0xFF, 0xFF, 0xFF), line=border, line_w=1.0,
         round_=True, adj=0.045)
    if title:
        txbox(slide, x + 0.18, y + 0.13, w - 0.36, 0.3, [[(title, 11.5, True, title_color)]])


def runs_label_body(label, body, lsz=9.5, bsz=9.0):
    return [[(label, lsz, True, DARK)], [(body, bsz, False, GRAY)]]


# =====================================================================
# Page 01 — 市场已成立，但还没有新的全民社交网络
# =====================================================================
s1 = prs.slides.add_slide(BLANK)
header(s1, 1, "市场已成立，但还没有新的全民社交网络",
       "收入与下载已具规模，但收入集中于少数产品，美国成人的陪伴用途仍属少数。")

AW, BW = 6.0, 6.0
AX, BX = MX, MX + AW + 0.23
AY, AH = CONTENT_Y, CONTENT_B - CONTENT_Y

card(s1, AX, AY, AW, AH, title="经营规模（可变现需求已出现）")
items_a = [
    ("移动变现｜2026Q1", "全球 AI Companion 移动 IAP 约 1.5 亿美元；季度收入较 2023Q1 超过 12 倍。"),
    ("获客规模｜2026Q1", "全球下载 8,300 万次，新增用户获取仍在扩张。"),
    ("收入集中｜2025 年中口径", "337 个仍活跃且产生收入的专门陪伴 App；头部 10% 获取约 89% 品类收入。"),
    ("成人用途｜美国，2026", "4% 美国成人用聊天机器人获取陪伴；10% 用于情绪支持或建议。"),
]
iy = AY + 0.52
step = (AH - 0.62) / 4
for i, (lab, body) in enumerate(items_a):
    badge(s1, AX + 0.18, iy + 0.03, i + 1)
    txbox(s1, AX + 0.60, iy, AW - 0.82, step - 0.10, runs_label_body(lab, body),
          space_after=1, line_spacing=1.02)
    iy += step

card(s1, BX, AY, BW, AH, title="普及边界（定义与样本）")
items_b = [
    ("核心讨论范围", "AI 原生关系、AI 内容社区与 AI 增强真人社交。"),
    ("外部替代与分发", "平台内嵌 AI 与免费捆绑；不计入核心陪伴市场规模。"),
]
iy = AY + 0.56
step = (AH - 0.70) / 2
for i, (lab, body) in enumerate(items_b):
    badge(s1, BX + 0.18, iy + 0.05, i + 5)
    txbox(s1, BX + 0.60, iy, BW - 0.82, step - 0.15, runs_label_body(lab, body),
          space_after=1, line_spacing=1.05)
    iy += step
txbox(s1, BX + 0.20, AY + AH - 0.72, BW - 0.4, 0.6,
      [[("口径提示：", 8.5, True, RED),
        ("不同调查的陪伴定义与样本不同，跨国、跨年份不可直接比较。", 8.5, False, GRAY)]],
      line_spacing=1.05)

banner(s1, "先评估具体关系场景与付费能力，再判断平台化空间。")
footer(s1, "来源：Sensor Tower、Appfigures、Pew、Common Sense Media；移动 IAP 与下载为估算口径，不含广告、Web、平台内嵌和硬件。")

# =====================================================================
# Page 02 — 中国玩家分争综合角色平台、垂直社区与真人社交增益
# =====================================================================
s2 = prs.slides.add_slide(BLANK)
header(s2, 2, "中国玩家分争综合角色平台、垂直社区与真人社交增益",
       "关系类型、内容供给与变现路径各不相同，应先选择竞争路线，再确定对标对象。")

ROWS_H, GAP = [1.98, 1.30, 1.30], 0.14
RY, _y = [], CONTENT_Y
for _h in ROWS_H:
    RY.append(_y); _y += _h + GAP

rows = [
    ("01", "综合角色平台",
     [("猫箱", "AI 角色、剧情与语音陪伴，体现产品化与付费能力。2026Q1 中国 iOS 生成式 AI 应用收入第 1；累计移动端消费者支出超 2,200 万美元（第三方估算）。"),
      ("星野 / Talkie", "UGC 角色与多模态互动，采用广告与订阅混合变现。2025 年前 9 个月合计平均 MAU 2,005.1 万、收入 1,875 万美元，广告约占 60%。")]),
    ("02", "垂直叙事 / IP 社区",
     [("筑梦岛", "女性向叙事与创作者 / IP 社区，依靠明确人群与内容供给定位。接近 500 万注册用户、约 80% 年轻女性、逾 50 万创作者（公司口径经媒体转引）。")]),
    ("03", "真人社交 AI 增益",
     [("Soul", "AI 为既有真人社交提供匹配与沟通增益，核心关系仍由真人承担。2025 年前 8 个月平均 DAU 1,100 万；AI Boosters 日均活跃 460 万（港交所材料）。")]),
]

for idx, (num, route, products) in enumerate(rows):
    y = RY[idx]
    ROW_H = ROWS_H[idx]
    rect(s2, MX, y, CW, ROW_H, fill=RGBColor(0xFF, 0xFF, 0xFF), line=BORDER, line_w=1.0,
         round_=True, adj=0.06)
    badge(s2, MX + 0.20, y + ROW_H / 2 - 0.15, num, d=0.30)
    txbox(s2, MX + 0.62, y + ROW_H / 2 - 0.16, 1.55, 0.36, [[(route, 11.5, True, DARK)]],
          anchor=MSO_ANCHOR.MIDDLE)
    rect(s2, MX + 2.28, y + 0.14, 0.016, ROW_H - 0.28, fill=RGBColor(0xD8, 0xCF, 0xCC))
    if len(products) == 2:
        sw_ = (CW - 2.55 - 0.24) / 2
        for j, (pname, pdesc) in enumerate(products):
            sx = MX + 2.55 + j * (sw_ + 0.24)
            rect(s2, sx, y + 0.14, sw_, ROW_H - 0.28, fill=RGBColor(0xFD, 0xF5, 0xF4),
                 line=BORDER, line_w=0.75, round_=True, adj=0.06)
            txbox(s2, sx + 0.16, y + 0.26, sw_ - 0.32, 0.28, [[(pname, 10.5, True, DARKRED)]])
            txbox(s2, sx + 0.16, y + 0.56, sw_ - 0.32, ROW_H - 0.82, [[(pdesc, 8.8, False, GRAY)]],
                  line_spacing=1.04)
    else:
        sx = MX + 2.55
        sw_ = CW - 2.55 - 0.22
        pname, pdesc = products[0]
        txbox(s2, sx, y + ROW_H / 2 - 0.15, 1.1, 0.3, [[(pname, 11, True, DARKRED)]])
        txbox(s2, sx + 1.15, y + ROW_H / 2 - 0.18, sw_ - 1.2, 0.4, [[(pdesc, 9.3, False, GRAY)]],
              anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.08)

banner(s2, "综合平台看关系留存，垂类社区看优质供给，真人社交看真实关系结果。")
footer(s2, "来源：Sensor Tower、Appfigures、36Kr、Soul 港交所材料；Soul 平台收入含 AI 收入；公司口径经媒体转引。")

# =====================================================================
# Page 03 — 海外竞争围绕四类优势展开，尚无稳定通吃者
# =====================================================================
s3 = prs.slides.add_slide(BLANK)
header(s3, 3, "海外竞争围绕四类优势展开，尚无稳定通吃者",
       "长期记忆、创作者供给、既有入口和真实关系，分别支撑不同产品路线；长期效果仍需检验。")

C1, C2, C3 = MX, MX + 1.95, MX + 8.35
CW1, CW2, CW3 = 1.32, 6.30, CW - 1.95 - 6.30
HY = CONTENT_Y
rect(s3, MX, HY, CW, 0.34, fill=RGBColor(0xF3, 0xEE, 0xEC), round_=True, adj=0.3)
txbox(s3, C1 + 0.03, HY + 0.03, CW1, 0.28, [[("优势类型", 9.5, True, GRAY)]],
      anchor=MSO_ANCHOR.MIDDLE)
txbox(s3, C2 + 0.10, HY + 0.03, CW2, 0.28, [[("代表产品 · 已有动作", 9.5, True, GRAY)]],
      anchor=MSO_ANCHOR.MIDDLE)
txbox(s3, C3 + 0.10, HY + 0.03, CW3 - 0.1, 0.28, [[("待验证", 9.5, True, GRAY)]],
      anchor=MSO_ANCHOR.MIDDLE)

rows3 = [
    ("01", "关系深度", "Replika / Nomi / Kindroid",
     "长期记忆、人格与持续陪伴，让价值从单次回答转向关系连续性。",
     "长期留存、心理安全与高推理成本能否兼容。"),
    ("02", "内容生态", "Character.AI；CHAI / PolyBuzz",
     "角色、场景、视频与 Remix / Feed 纳入变现；CHAI 与 PolyBuzz 推进移动端广告和订阅。CHAI 自报实验：广告由每 8 条消息一次降至每 16 条一次，D30 留存相对提升 18.5%。",
     "观看和创作能否形成稳定网络；变现强度是否损害连续互动。"),
    ("03", "平台分发", "Meta AI / AI Studio",
     "把 AI 人格嵌入既有社交产品与独立 App，以社交图谱和跨产品入口降低获客摩擦。",
     "平台品牌与监管边界下，深度角色体验能否持续。"),
    ("04", "真人关系", "Tinder / Bumble",
     "推荐、Chemistry 与 AI Photo Feedback 辅助匹配和表达，最终关系仍由真人承担。",
     "AI 增益如何归因，并守住身份真实性与隐私边界。"),
]
ry = HY + 0.42
rh = (CONTENT_B - ry - 0.06) / 4
for num, adv, prods, action, verify in rows3:
    rect(s3, MX, ry, CW, rh, fill=RGBColor(0xFF, 0xFF, 0xFF), line=BORDER, line_w=0.75,
         round_=True, adj=0.05)
    badge(s3, C1 + 0.02, ry + rh / 2 - 0.15, num, d=0.30)
    txbox(s3, C1 + 0.42, ry + 0.10, CW1 - 0.45, rh - 0.2, [[(adv, 10.5, True, DARK)]],
          anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.02)
    txbox(s3, C2, ry + 0.09, CW2, 0.28, [[(prods, 9.5, True, DARKRED)]])
    txbox(s3, C2, ry + 0.36, CW2, rh - 0.45, [[(action, 8.8, False, GRAY)]],
          line_spacing=1.03)
    txbox(s3, C3, ry, CW3, rh,
          [[("待验证｜", 9.0, True, RED), (verify, 9.0, False, GRAY)]],
          anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.04)
    ry += rh + 0.045
    if num != "04":
        hline(s3, MX + 0.1, ry - 0.035, CW - 0.2, color=RGBColor(0xEE, 0xE6, 0xE3))

banner(s3, "独立产品要验证自身优势，能否抵御平台分发与免费能力的挤压。")
footer(s3, "来源：原研究所列产品官方披露；截至 2026-08-09。功能上线非网络效应；CHAI 为特定公司实验，非行业结论；四类优势不互斥。")

# =====================================================================
# Page 04 — 未来增长取决于关系价值、单位经济与安全能否同时成立
# =====================================================================
s4 = prs.slides.add_slide(BLANK)
header(s4, 4, "未来增长取决于关系价值、单位经济与安全能否同时成立",
       "三种增长情景取决于不同条件，长期经营仍需同时验证关系连续性、内容供给、盈利与治理。")

# --- 左栏：2026 锚点 ---
LX, LW = MX, 2.75
txbox(s4, LX, CONTENT_Y, LW, 0.3, [[("2026 锚点（历史锚点）", 11, True, DARK)]])
circ = s4.shapes.add_shape(MSO_SHAPE.OVAL, Inches(LX + 0.55), Inches(CONTENT_Y + 0.42),
                           Inches(1.65), Inches(1.65))
circ.fill.solid(); circ.fill.fore_color.rgb = RED
circ.line.color.rgb = DARKRED; circ.line.width = Pt(1.25); circ.shadow.inherit = False
tf = circ.text_frame; tf.word_wrap = True
tf.margin_left = tf.margin_right = Inches(0.08)
tf.margin_top = tf.margin_bottom = Inches(0.05)
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r1 = p.add_run(); r1.text = "6 亿美元"; set_font(r1, 15, True, RGBColor(0xFF, 0xFF, 0xFF))
p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
r2 = p2.add_run(); r2.text = "2026 年移动 IAP 运行率"; set_font(r2, 8.5, False, RGBColor(0xFF, 0xE3, 0xE0))
txbox(s4, LX - 0.1, CONTENT_Y + 2.2, LW + 0.2, 0.6,
      [[("锚点：2026Q1 移动 IAP 1.5 亿美元（简单年化）", 8.8, False, GRAY)],
       [("历史锚点 → 假设推演（非预测）", 8.8, False, GRAY)]], space_after=2, line_spacing=1.05)

# --- 中栏：三情景 ---
MX2, MW2 = LX + LW + 0.28, 5.85
txbox(s4, MX2, CONTENT_Y, MW2, 0.3, [[("三种情景（条件不同 → 结果不同）", 11, True, DARK)]])
scen = [
    ("下行", "假设 CAGR 10%", "未成年人口收紧、低留存及平台免费捆绑", "8.0"),
    ("基准", "假设 CAGR 30%", "记忆、语音、角色生态与多元变现改善增长质量", "13.2"),
    ("上行", "假设 CAGR 50%", "多模态、创作者网络效应与可规模化合规共同成立", "20.3"),
]
sy = CONTENT_Y + 0.40
sh_, sgap = 1.25, 0.16
for i, (tag, cagr, cond, val) in enumerate(scen):
    y = sy + i * (sh_ + sgap)
    rect(s4, MX2, y, MW2, sh_, fill=RGBColor(0xFF, 0xFF, 0xFF), line=BORDER, line_w=1.0,
         round_=True, adj=0.07)
    rect(s4, MX2, y, 0.07, sh_, fill=RED if tag != "基准" else DARKRED)
    txbox(s4, MX2 + 0.22, y + 0.20, 1.5, 0.3, [[(tag + "｜", 10.5, True, RED), (cagr, 10.5, True, DARK)]])
    txbox(s4, MX2 + 0.22, y + 0.52, MW2 - 1.9, sh_ - 0.62, [[(cond, 8.8, False, GRAY)]],
          line_spacing=1.05)
    txbox(s4, MX2 + MW2 - 2.0, y + 0.16, 1.85, sh_ - 0.3,
          [[("2029 年约", 9, False, GRAY)], [(val + " 亿美元", 15, True, DARK)]],
          align=PP_ALIGN.RIGHT, space_after=0, line_spacing=1.0)
    # 连接线：锚点圆 → 情景卡
    ycirc = CONTENT_Y + 0.42 + 1.65 / 2
    ln = s4.shapes.add_connector(2, Inches(LX + 2.35), Inches(ycirc), Inches(MX2), Inches(y + sh_ / 2))
    ln.line.color.rgb = RED; ln.line.width = Pt(1.2)

# --- 右栏：经营跟踪指标 ---
RX, RW = MX2 + MW2 + 0.28, SW - MX - (MX2 + MW2 + 0.28)
txbox(s4, RX, CONTENT_Y, RW, 0.3, [[("经营跟踪指标（持续观察，尚待验证）", 10.5, True, DARK)]])
kpis = [
    ("关系连续性", "30 / 90 日关系留存；记忆命中率与人格稳定性。"),
    ("内容供给", "优质角色、活跃创作者与分成。"),
    ("单位经济", "推理 / 审核成本、贡献毛利与退款率。"),
    ("关系治理", "年龄识别、危机干预与强化儿童安全、地区规则覆盖。"),
]
kgy = CONTENT_Y + 0.42
kgh = 1.80                                   # 两行卡高
kgw = (RW - 0.12) / 2                        # 两列
for i, (k, v) in enumerate(kpis):
    xx = RX + (i // 2) * (kgw + 0.12)
    yy = kgy + (i % 2) * (kgh + 0.10)
    rect(s4, xx, yy, kgw, kgh, fill=RGBColor(0xFD, 0xF5, 0xF4), line=BORDER, line_w=0.75,
         round_=True, adj=0.09)
    rect(s4, xx, yy, 0.05, kgh, fill=RED)
    txbox(s4, xx + 0.16, yy + 0.22, kgw - 0.30, 0.26, [[(k, 9.5, True, DARK)]])
    txbox(s4, xx + 0.16, yy + 0.50, kgw - 0.30, kgh - 0.62, [[(v, 8.2, False, GRAY)]],
          line_spacing=1.06)
txbox(s4, RX, kgy + 2 * kgh + 0.22, RW, 0.5,
      [[("注：", 8, True, RED), ("以上指标均为持续观察项，尚未验证；需随监管落地与季报数据持续校准。", 8, False, GRAY)]],
      line_spacing=1.08)

banner(s4, "用「安全关系留存 × 每用户贡献毛利」检验增长质量。")
footer(s4, "来源：原研究测算与监管材料；截至 2026-08-09。三年 CAGR 为假设、置信度中低；仅移动 IAP，非完整 TAM；6 亿美元为简单年化运行率，非全年业绩。")

# ---------- 保存 ----------
OUT = "workspace/AI社交软件市场竞争格局与未来展望.pptx"
prs.save(OUT)
print("saved:", OUT)
